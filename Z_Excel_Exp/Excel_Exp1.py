import openpyxl
import time

def append_data_to_excel(filename, data, sheet_name="Sheet1"):
  
  try:
    # Load the existing workbook
    wb = openpyxl.load_workbook(filename)
    sheet = wb[sheet_name]
    last_row = sheet.max_row + 1
  except FileNotFoundError:
    # Create a new workbook and sheet if the file doesn't exist
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = sheet_name

    # Write the header row
    sheet['A1'] = 'Article'
    sheet['B1'] = 'Para Number'
    sheet['C1'] = 'Para Content'
    sheet['D1'] = 'Software Requirement'

    # Get the last row number
    last_row = sheet.max_row

  # Append data to the sheet
  for row_data in data:
    last_row += 1
    for col_num, value in enumerate(row_data.values(), 1):
      cell = sheet.cell(row=last_row, column=col_num)
      cell.value = value

  # Save the workbook
  wb.save(filename)

file = "./Z_Excel_Exp/CELEX_02019R0817-20240425_EN_V1.xlsx"
result_data = []
result = {"Article":'Key1',"Para Number":"para_num 123","Para Content":"Test Data","Software Requirement":"No"} 
result_data.append(result)

append_data_to_excel(file, result_data)